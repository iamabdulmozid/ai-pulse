import json
import logging

from django.conf import settings
from django.contrib.auth.decorators import login_required
from django.http import HttpResponse, JsonResponse, StreamingHttpResponse
from django.shortcuts import render
from django.views.decorators.csrf import csrf_protect

from services import metrics
from services.assistant import router

logger = logging.getLogger(__name__)


def _resolve(run, message, user):
    """Use OpenAI tool-calling when a key is configured; fall back to the deterministic router on any
    error or when no key is set (keeps the demo working offline)."""
    use_llm = bool(settings.OPENAI_API_KEY) and not settings.ASSISTANT_FALLBACK_MODE
    if use_llm:
        try:
            from services.assistant.llm import answer_llm

            result = answer_llm(run, message, user)
            if result.get("text"):
                result["engine"] = "openai"
                return result
        except Exception:  # noqa: BLE001
            logger.exception("OpenAI assistant path failed; using deterministic fallback")
    result = router.answer(run, message)
    result["engine"] = "fallback"
    return result


@login_required
def assistant_page(request):
    return render(request, "assistant/page.html", {"thread": None})


@login_required
def assistant_panel(request):
    return render(request, "assistant/_slideover.html")


def _sse(event: str, data) -> str:
    return f"event: {event}\ndata: {json.dumps(data)}\n\n"


@login_required
@csrf_protect
def assistant_stream(request):
    message = (request.POST.get("message") or request.GET.get("message") or "").strip()
    run = metrics.latest_run()

    def gen():
        if run is None:
            yield _sse("error", {"message": "No data yet."})
            yield _sse("done", {})
            return
        result = _resolve(run, message, request.user)
        for i, label in enumerate(result.get("steps", [])):
            yield _sse("step", {"id": i, "label": label, "status": "done"})
        yield _sse("delta", {"text": result["text"]})
        if result.get("table"):
            yield _sse("table", result["table"])
        if result.get("chart"):
            yield _sse("chart", result["chart"])
        yield _sse("sources", result.get("sources", []))
        if result.get("followups"):
            yield _sse("followups", result["followups"])
        _persist(request.user, message, result)
        yield _sse("done", {})

    resp = StreamingHttpResponse(gen(), content_type="text/event-stream")
    resp["Cache-Control"] = "no-cache"
    resp["X-Accel-Buffering"] = "no"
    return resp


def _persist(user, message, result):
    try:
        from apps.assistant.models import ChatMessage, Conversation, TokenUsage

        conv = Conversation.objects.create(user=user, title=message[:120] or "Untitled")
        ChatMessage.objects.create(conversation=conv, role="user", content=message)
        ChatMessage.objects.create(
            conversation=conv, role="assistant", content=result["text"],
            artifacts={"table": result.get("table"), "chart": bool(result.get("chart"))},
            sources=result.get("sources"),
        )
        TokenUsage.objects.create(conversation=conv, model="fallback", total_tokens=0)
    except Exception:
        pass


@login_required
def threads_list(request):
    from apps.assistant.models import Conversation

    rows = list(Conversation.objects.filter(user=request.user).values("id", "title", "updated_at")[:30])
    return JsonResponse({"threads": [{"id": str(r["id"]), "title": r["title"]} for r in rows]})


@login_required
def assistant_export(request):
    """Export a returned table to xlsx."""
    import openpyxl
    from openpyxl.utils import get_column_letter

    payload = json.loads(request.body or "{}")
    cols = payload.get("columns", [])
    rows = payload.get("rows", [])
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Answer"
    ws.append(cols)
    for r in rows:
        ws.append(r)
    for i in range(len(cols)):
        ws.column_dimensions[get_column_letter(i + 1)].width = 22
    resp = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp["Content-Disposition"] = "attachment; filename=ai-pulse-answer.xlsx"
    wb.save(resp)
    return resp
