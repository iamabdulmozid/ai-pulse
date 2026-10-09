"""Generate the AI Pulse demo data set.

Writes six Excel files that a Karbar merchandiser would upload, two extra
daily reports for the live upload demo, and an answer key that holds the
expected predictions for every open PO.

Everything is deterministic: the same SEED always produces the same files.
The reference prediction engine lives in this file too, so the answer key and
the demo story ("86% on time, USD 2.4M at risk, hero PO 9 days late") are
computed, not typed in.

    python scripts/generate_sample_data.py
"""

from __future__ import annotations

import math
import random
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, timedelta
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

SEED = 20261015
TODAY = date(2026, 10, 15)                 # demo day, a Thursday
LAST_EXPECTED_REPORT = date(2026, 10, 14)  # factories report the previous working day
HISTORY_START = date(2025, 10, 15)         # 12 months of shipped history
DPR_START = date(2026, 9, 1)               # daily production rows start here
OUT = Path(__file__).resolve().parent.parent / "sample_data"
ONE = timedelta(days=1)

# --- Reference engine parameters (mirrored in README.md and the answer key) ---
AIR_USD_PER_KG = 6.00
SEA_USD_PER_KG = 0.50
KNIT_MC_MINUTES_PER_DAY = 1020          # 2 shifts x 10 h x 85% efficiency
RATE_WINDOW = 7                         # working days in the weighted average
YARN_TO_KNIT_DAYS = 2                   # knitting can start 2 days after yarn in-house
FLOW_LIMIT_DAYS = 1.5                   # work-in-front <= 1.5 days of output: the stage is supply-limited
STAGE_LAG = {"linking": 1.0, "mending": 0.5, "washing": 0.5, "ironing": 0.5, "packing": 0.5}
ALL_STAGES = ["knitting", "linking", "mending", "washing", "ironing", "packing"]
STAGE_LABEL = {
    "knitting": "Knitting", "linking": "Linking", "mending": "Trimming & Mending",
    "washing": "Washing", "ironing": "Ironing", "packing": "Packing",
}
RAMP = [0.5, 0.75, 0.9]

# Targets for the demo story
TARGET_OPEN_FOB = 18_600_000
TARGET_OPEN_PCS = 1_900_000
TARGET_VAR = 2_400_000
TARGET_VAR_RISKY_SHARE = 0.70

rng = random.Random(SEED)


# ---------------------------------------------------------------- calendar ---
def is_wd(d: date, extra: frozenset = frozenset()) -> bool:
    """Friday is the weekend in Bangladesh garment factories."""
    return d.weekday() != 4 or d in extra


def wd_between(a: date, b: date, extra: frozenset = frozenset()) -> int:
    """Working days d with a <= d < b."""
    n, d = 0, a
    while d < b:
        n += is_wd(d, extra)
        d += ONE
    return n


def nth_wd(start: date, n: int, extra: frozenset = frozenset()) -> date:
    """The n-th working day counting `start` as day 1 when it is a working day."""
    d, c = start, 0
    while True:
        if is_wd(d, extra):
            c += 1
            if c == n:
                return d
        d += ONE


def next_wd(d: date) -> date:
    d += ONE
    while not is_wd(d):
        d += ONE
    return d


def prev_wd(d: date) -> date:
    d -= ONE
    while not is_wd(d):
        d -= ONE
    return d


def wd_list(a: date, b: date) -> list[date]:
    """Working days a..b inclusive."""
    out, d = [], a
    while d <= b:
        if is_wd(d):
            out.append(d)
        d += ONE
    return out


def not_friday(d: date) -> date:
    return d - ONE if d.weekday() == 4 else d


def rand_date(a: date, b: date, r: random.Random) -> date:
    return not_friday(a + timedelta(days=r.randint(0, (b - a).days)))


# --------------------------------------------------------------- master data -
@dataclass
class Factory:
    code: str
    name: str
    area: str
    district: str
    gauges: tuple
    weight: float
    otd_target: float
    last_report: date
    merch: str
    risky: bool = False


MERCHANDISERS = ["Farhana Rahman", "Tanvir Hasan", "Nusrat Jahan", "Arif Chowdhury",
                 "Sadia Islam", "Mehedi Karim", "Rumana Akter", "Imran Hossain"]
INSPECTORS = ["Kamrul Hasan", "Shirin Sultana", "Rakib Uddin", "Lipi Das", "Jahid Alam", "Moumita Saha"]

F = lambda *a, **k: Factory(*a, **k)  # noqa: E731
LR = LAST_EXPECTED_REPORT
FACTORIES = [
    F("GRL", "Greyloom Knitwear Ltd.", "Konabari", "Gazipur", (5, 7, 12), 9.0, 0.68, LR, MERCHANDISERS[0], True),
    F("IRB", "Ironbridge Sweaters Ltd.", "Fatullah", "Narayanganj", (3, 5, 7), 7.5, 0.70, LR, MERCHANDISERS[1], True),
    F("SLM", "Saltmarsh Knit Fashions Ltd.", "Ashulia", "Dhaka", (7, 12, 14), 6.5, 0.72, date(2026, 10, 12), MERCHANDISERS[2], True),
    F("KST", "Kestrel Knitwear Ltd.", "Kashimpur", "Gazipur", (7, 12, 14), 5.5, 0.93, LR, MERCHANDISERS[3]),
    F("BLF", "Bluefinch Sweaters Ltd.", "Zirabo", "Dhaka", (3, 5, 7), 4.5, 0.90, LR, MERCHANDISERS[4]),
    F("CPL", "Copperleaf Knit Ltd.", "DEPZ, Savar", "Dhaka", (12, 14), 4.5, 0.95, LR, MERCHANDISERS[5]),
    F("TLG", "Tallgrass Knitwear Ltd.", "Sreepur", "Gazipur", (5, 7, 12), 5.0, 0.88, LR, MERCHANDISERS[6]),
    F("RVS", "Riverstone Sweaters Ltd.", "Siddhirganj", "Narayanganj", (3, 5, 7, 12), 4.5, 0.86, LR, MERCHANDISERS[7]),
    F("AMB", "Amberline Knit Ltd.", "CEPZ", "Chattogram", (7, 12, 14), 4.0, 0.94, LR, MERCHANDISERS[3]),
    F("SVP", "Silverpine Knitwear Ltd.", "Tongi", "Gazipur", (12, 14), 3.5, 0.92, LR, MERCHANDISERS[5]),
    F("LTF", "Lanternfield Sweaters Ltd.", "Mawna", "Gazipur", (5, 7), 3.5, 0.87, LR, MERCHANDISERS[6]),
    F("MDW", "Morningdew Knit Ltd.", "Ashulia", "Dhaka", (7, 12), 4.0, 0.91, LR, MERCHANDISERS[2]),
    F("OKH", "Oakhaven Knitwear Ltd.", "Kalurghat", "Chattogram", (7, 12, 14), 3.0, 0.89, date(2026, 10, 11), MERCHANDISERS[4]),
    F("WTH", "Wildthorn Sweaters Ltd.", "Kanchpur", "Narayanganj", (3, 5, 7), 3.5, 0.85, LR, MERCHANDISERS[1]),
    F("BRW", "Brightwool Knit Industries Ltd.", "Kashimpur", "Gazipur", (5, 7, 12), 4.5, 0.90, LR, MERCHANDISERS[0]),
    F("FNS", "Fernstitch Knitwear Ltd.", "Savar", "Dhaka", (12, 14), 3.0, 0.96, LR, MERCHANDISERS[5]),
    F("HMR", "Highmoor Sweaters Ltd.", "Bhaluka", "Mymensingh", (3, 5, 7), 3.0, 0.84, date(2026, 10, 13), MERCHANDISERS[7]),
    F("CLB", "Cloudberry Knit Ltd.", "Chandra", "Gazipur", (7, 12), 3.5, 0.92, LR, MERCHANDISERS[6]),
    F("SND", "Sundial Knitwear Ltd.", "Baipail, Ashulia", "Dhaka", (7, 12, 14), 3.5, 0.93, LR, MERCHANDISERS[2]),
    F("RSG", "Rosegate Sweaters Ltd.", "Fatullah", "Narayanganj", (3, 5), 2.5, 0.83, LR, MERCHANDISERS[1]),
    F("PBB", "Pebblebrook Knit Ltd.", "CEPZ", "Chattogram", (12, 14), 3.0, 0.97, date(2026, 10, 13), MERCHANDISERS[4]),
    F("WSM", "Westmere Knitwear Ltd.", "Tongi", "Gazipur", (5, 7, 12), 3.5, 0.89, LR, MERCHANDISERS[7]),
]
FAC = {f.code: f for f in FACTORIES}
RISKY = [f.code for f in FACTORIES if f.risky]

# product: dept, gauges, linking pcs/machine/day, knit-minute multiplier, weight multiplier, fob multiplier
PRODUCTS = [
    ("Men", "Crew-neck Pullover", (5, 7, 12), 36, 1.00, 1.00, 1.00),
    ("Men", "V-neck Pullover", (7, 12), 36, 1.00, 0.97, 1.00),
    ("Men", "Half-zip Pullover", (5, 7), 30, 1.10, 1.05, 1.08),
    ("Men", "Shawl-collar Cardigan", (3, 5), 28, 1.30, 1.15, 1.15),
    ("Men", "Knitted Vest", (12,), 48, 0.70, 0.70, 0.80),
    ("Women", "Crew-neck Pullover", (7, 12, 14), 36, 1.00, 0.90, 1.00),
    ("Women", "V-neck Cardigan", (7, 12), 32, 1.20, 1.00, 1.10),
    ("Women", "Turtle-neck Pullover", (7, 12), 34, 1.10, 0.95, 1.05),
    ("Women", "Cropped Pullover", (5, 7), 38, 0.85, 0.80, 0.95),
    ("Women", "Long Cardigan", (3, 5), 26, 1.45, 1.25, 1.20),
    ("Kids", "Crew-neck Jumper", (7, 12), 50, 0.60, 0.50, 0.60),
    ("Kids", "Cardigan", (7,), 46, 0.70, 0.55, 0.65),
    ("Kids", "Jacquard Jumper", (7,), 44, 0.85, 0.55, 0.70),
]
GAUGE_KNIT_MIN = {3: 22, 5: 30, 7: 42, 12: 72, 14: 88}
GAUGE_WEIGHT = {3: 0.85, 5: 0.75, 7: 0.62, 12: 0.42, 14: 0.35}
GAUGE_FOB = {3: 1.15, 5: 1.05, 7: 1.00, 12: 1.00, 14: 1.05}
PCS_PER_CARTON = {3: 12, 5: 12, 7: 18, 12: 24, 14: 24}
# yarn: (name, short name, fob range, wash, weight mult, gauges)
YARNS_AW = [
    ("100% Lambswool", "Lambswool", (14.0, 18.0), "Y", 1.10, (3, 5, 7)),
    ("100% Extra-fine Merino Wool", "Merino", (16.0, 21.0), "Y", 0.95, (7, 12, 14)),
    ("70% Acrylic 30% Wool", "Wool-blend", (9.0, 12.0), "Y", 1.00, (3, 5, 7, 12)),
    ("100% Acrylic", "Acrylic", (6.5, 9.0), "N", 0.90, (3, 5, 7, 12)),
    ("60% Cotton 40% Acrylic", "Cotton-blend", (8.0, 10.5), "N", 1.00, (5, 7, 12, 14)),
    ("50% Recycled Polyester 50% Acrylic", "Recycled", (7.0, 9.0), "N", 0.90, (5, 7, 12)),
]
YARNS_SS = [
    ("100% Cotton", "Cotton", (8.5, 11.5), "N", 1.05, (12, 14)),
    ("100% Organic Cotton", "Organic Cotton", (9.5, 12.5), "N", 1.05, (12, 14)),
    ("55% Linen 45% Cotton", "Linen-blend", (11.0, 14.0), "Y", 1.00, (12, 14)),
    ("60% Cotton 40% Acrylic", "Cotton-blend", (8.0, 10.5), "N", 1.00, (12, 14)),
]
COLOURS_AW = ["Navy", "Charcoal Melange", "Oatmeal Melange", "Camel", "Burgundy", "Forest Green",
              "Ecru", "Black", "Grey Marl", "Rust", "Cream", "Olive"]
COLOURS_SS = ["Ecru", "Sky Blue", "White", "Sage", "Dusty Pink", "Navy", "Butter Yellow", "Stone"]
SIZES = {
    "Men": (["S", "M", "L", "XL", "XXL"], [15, 30, 30, 17, 8]),
    "Women": (["XS", "S", "M", "L", "XL"], [10, 25, 32, 22, 11]),
    "Kids": (["2-3Y", "4-5Y", "6-7Y", "8-9Y", "10-11Y"], [15, 22, 25, 22, 16]),
}
ALL_SIZES = ["XS", "S", "M", "L", "XL", "XXL", "2-3Y", "4-5Y", "6-7Y", "8-9Y", "10-11Y"]
DESTINATIONS = [("Hamburg, Germany", 0.30, (30, 38)), ("Rotterdam, Netherlands", 0.20, (28, 36)),
                ("Felixstowe, United Kingdom", 0.20, (32, 40)), ("Newark, United States", 0.18, (35, 45)),
                ("Los Angeles, United States", 0.12, (28, 35))]
VESSELS = ["MV Bay Meridian", "MV Coral Horizon", "MV Delta Pioneer", "MV Silver Monsoon",
           "MV Northern Lotus", "MV Eastern Kestrel", "MV Jade Estuary", "MV Amber Tide"]
DEFECTS = ["Hole", "Dropped stitch", "Uneven linking", "Measurement out of tolerance",
           "Shading between panels", "Pilling", "Loose yarn ends", "Weight below tolerance",
           "Mismatched stripes", "Oil stain"]

# T&A template: (seq, milestone, days before ex-factory, responsible)
TA_TEMPLATE = [
    (1, "Yarn booking", 95, "Factory"),
    (2, "Yarn shade approval", 85, "Karbar Merchandising"),
    (3, "Fit sample approval", 80, "Karbar Merchandising"),
    (4, "Size set approval", 60, "Karbar Merchandising"),
    (5, "PP sample approval", 50, "Karbar Merchandising"),
    (6, "Yarn in-house", 45, "Factory / Yarn supplier"),
    (7, "PP meeting", 42, "Factory + Karbar QA"),
    (8, "Knitting start", 40, "Factory"),
    (9, "Linking start", 34, "Factory"),
    (10, "Final inspection", 1, "Karbar QA"),
    (11, "Ex-factory", 0, "Factory"),
]


def season_for(exf: date) -> str:
    y = exf.year % 100
    if exf.month >= 12 and exf.day >= 21:
        return f"SS{y + 1}"
    if exf.month in (1, 2, 3):
        return f"SS{y}"
    if exf.month in (4, 5, 6):
        return f"SU{y}"
    if exf.month == 12 or (exf.month == 11 and exf.day >= 16):
        return f"HO{y}"
    return f"AW{y}"


def split_qty(total: int, ratios: list[float]) -> list[int]:
    """Largest-remainder split so the parts always sum to `total`."""
    s = sum(ratios)
    raw = [total * r / s for r in ratios]
    base = [int(x) for x in raw]
    order = sorted(range(len(raw)), key=lambda i: raw[i] - base[i], reverse=True)
    for i in order[: total - sum(base)]:
        base[i] += 1
    return base


# ------------------------------------------------------------------- orders --
@dataclass
class PO:
    po_no: str
    po_date: date
    season: str
    dept: str
    product: str
    gauge: int
    yarn: str
    yarn_short: str
    wash: str
    weight: float
    knit_min: float
    link_std: int
    factory: str
    qty: int
    fob: float
    exf: date
    destination: str
    style_no: str = ""
    style_name: str = ""
    merch: str = ""
    open: bool = True
    scenario: str = "ok"          # ok, watch, tight, pred_late, late, hero, donor, closed
    target_slack: int | None = None
    target_late_wd: int | None = None
    yarn_delay: int = 0
    pp_delay: int = 0
    yarn_revised: date | None = None
    actual_exf: date | None = None
    knit_start: date | None = None
    link_start: date | None = None
    rows: list = field(default_factory=list)       # simulated daily rows (full, to Oct 14)
    colours: list = field(default_factory=list)

    @property
    def value(self) -> float:
        return round(self.qty * self.fob, 2)

    @property
    def stages(self) -> list[str]:
        return [s for s in ALL_STAGES if s != "washing" or self.wash == "Y"]

    @property
    def planned_knit(self) -> date:
        return not_friday(self.exf - timedelta(days=40))

    @property
    def planned_link(self) -> date:
        return not_friday(self.exf - timedelta(days=34))

    @property
    def planned_yarn(self) -> date:
        return not_friday(self.exf - timedelta(days=45))


def pick_product(fac: Factory, season: str, r: random.Random):
    ss = season.startswith("SS") or season.startswith("SU")
    allowed = set(fac.gauges) & ({12, 14} if ss else {3, 5, 7, 12})
    if not allowed:
        allowed = set(fac.gauges)
    for _ in range(50):
        dept = r.choices(["Men", "Women", "Kids"], [0.4, 0.4, 0.2])[0]
        cands = [p for p in PRODUCTS if p[0] == dept and set(p[2]) & allowed]
        if cands:
            p = r.choice(cands)
            gauge = r.choice(sorted(set(p[2]) & allowed))
            yarns = [y for y in (YARNS_SS if ss else YARNS_AW) if gauge in y[5]] or (YARNS_SS if ss else YARNS_AW)
            return p, gauge, r.choice(yarns)
    raise RuntimeError("no product for factory")


def build_po(po_no, exf, fac: Factory, r: random.Random, open_: bool, product=None, qty=None, fob=None) -> PO:
    season = season_for(exf)
    if product is None:
        p, gauge, y = pick_product(fac, season, r)
    else:
        p, gauge, y = product
    dept, name, _, link_std, km, wm, fm = p
    weight = round(GAUGE_WEIGHT[gauge] * wm * y[4] * (0.55 if dept == "Kids" else 1.0), 2)
    knit_min = round(GAUGE_KNIT_MIN[gauge] * km, 1)
    if fob is None:
        fob = round(r.uniform(*y[2]) * GAUGE_FOB[gauge] * fm, 2)
    if qty is None:
        q = int(r.lognormvariate(math.log(4000), 0.5))
        qty = max(1200, min(15000, q)) // 12 * 12
    destination = r.choices([d[0] for d in DESTINATIONS], [d[1] for d in DESTINATIONS])[0]
    po = PO(po_no, exf - timedelta(days=r.randint(120, 160)), season, dept, name, gauge, y[0], y[1],
            y[3], weight, knit_min, link_std, fac.code, qty, fob, exf, destination, open=open_)
    po.merch = fac.merch if r.random() < 0.85 else r.choice(MERCHANDISERS)
    return po


def assign_styles(pos: list[PO], r: random.Random):
    used = set()
    for po in pos:
        while True:
            sn = f"K{po.season}-{po.dept[0]}-{r.randint(1000, 9999)}"
            if sn not in used:
                used.add(sn)
                break
        po.style_no = sn
        po.style_name = f"{po.yarn_short} {po.product}" if po.dept != "Kids" else f"Kids {po.yarn_short} {po.product}"
        palette = COLOURS_SS if po.season.startswith(("SS", "SU")) else COLOURS_AW
        n = r.choice([1, 2, 2, 3, 3, 3, 4]) if po.qty >= 2400 else r.choice([1, 2])
        cols = r.sample(palette, n)
        shares = [r.uniform(0.7, 1.3) for _ in cols]
        po.colours = list(zip(cols, split_qty(po.qty, shares)))


# --------------------------------------------------------------- simulation --
def simulate(po: PO, base: float, start: date, end: date, mult: dict, link_delay: int,
             seed: int, forced_link: dict | None = None, until_packed: bool = False) -> list[dict]:
    """Day-by-day production. Each stage draws only on the previous day's upstream output."""
    r = random.Random(seed)
    stages = po.stages
    days = wd_list(start, end)
    noise = {s: [r.uniform(0.92, 1.08) for _ in days] for s in stages}
    sidx = {stages[0]: 0}
    for i, s in enumerate(stages[1:], 1):
        sidx[s] = sidx[stages[i - 1]] + (link_delay if s == "linking" else 1)
    if forced_link:
        first = min(forced_link)
        sidx["linking"] = days.index(first) if first in days else len(days)
    cum = {s: 0 for s in stages}
    rows = []
    for i, d in enumerate(days):
        prev = dict(cum)
        out = {}
        for j, s in enumerate(stages):
            if i < sidx[s]:
                out[s] = 0
                continue
            avail = po.qty - cum[s] if j == 0 else prev[stages[j - 1]] - cum[s]
            if s == "linking" and forced_link is not None:
                o = forced_link.get(d, 0)
            else:
                k = i - sidx[s]
                o = int(round(base * mult[s] * (RAMP[k] if k < 3 else 1.0) * noise[s][i]))
            out[s] = max(0, min(o, avail))
        for s in stages:
            cum[s] += out[s]
        k_k = i - sidx["knitting"]
        k_l = i - sidx["linking"]
        kmc = round(base * mult["knitting"] * (RAMP[k_k] if 0 <= k_k < 3 else 1) / (KNIT_MC_MINUTES_PER_DAY / po.knit_min)) if out["knitting"] else 0
        lmc = round(base * mult["linking"] * (RAMP[k_l] if 0 <= k_l < 3 else 1) / po.link_std) if out["linking"] else 0
        rows.append({"date": d, "out": out, "cum": dict(cum), "kmc": max(kmc, 1 if out["knitting"] else 0),
                     "lmc": max(lmc, 1 if out["linking"] else 0)})
        if until_packed and cum["packing"] >= po.qty:
            break
        if cum["packing"] >= po.qty:
            break
    return rows


# -------------------------------------------------------- reference engine ---
def weighted_rate(values: list[int]) -> float:
    if not values:
        return 0.0
    w = range(1, len(values) + 1)
    return sum(a * b for a, b in zip(w, values)) / sum(w)


def stage_rates(po: PO, rows: list[dict], last_report: date) -> dict:
    window = []
    d = last_report
    while len(window) < RATE_WINDOW:
        if is_wd(d):
            window.append(d)
        d -= ONE
    window.reverse()
    by_date = {x["date"]: x for x in rows}
    rates = {}
    for s in po.stages:
        firsts = [x["date"] for x in rows if x["out"][s] > 0]
        if not firsts:
            rates[s] = 0.0
            continue
        days = [d for d in window if d >= firsts[0]]
        rates[s] = weighted_rate([by_date[d]["out"][s] if d in by_date else 0 for d in days])
    return rates


def project(po: PO, rows: list[dict], last_report: date, extra: frozenset = frozenset(),
            link_mult: float = 1.0) -> dict:
    """Projection for a PO with knitting under way. `rows` are reported rows only."""
    stages = po.stages
    cum = rows[-1]["cum"] if rows else {s: 0 for s in stages}
    meas = stage_rates(po, rows, last_report)
    meas["linking"] *= link_mult
    info, prev, f_prev, rem_prev, eff_prev = {}, None, 0.0, 0, None
    for s in stages:
        rem = po.qty - cum[s]
        m = meas[s]
        # A stage that has not started, or whose work-in-front is small, is limited by supply:
        # its measured rate shows what it was fed, not what it can do, so it follows upstream.
        # Once upstream is finished the stage works through what is left at its own pace.
        limited = prev is not None and (cum[s] == 0 or cum[prev] - cum[s] <= FLOW_LIMIT_DAYS * m)
        follows = limited and rem_prev > 0
        eff = max(m, eff_prev) if limited else m
        lower = f_prev + STAGE_LAG[s] if (prev is not None and rem_prev > 0) else 0.0
        if rem <= 0:
            fin, binding = 0.0, "done"
        elif follows:
            fin, binding = lower, "flow"
        else:
            own = rem / eff if eff > 0 else math.inf
            fin, binding = (own, "own") if own >= lower else (lower, "flow")
        info[s] = {"cum": cum[s], "rem": rem, "rate": m, "eff": eff, "finish": fin, "binding": binding}
        prev, f_prev, rem_prev, eff_prev = s, fin, rem, eff
    finish = f_prev
    n = 0 if finish <= 0 else math.ceil(finish - 1e-9)
    avail = wd_between(TODAY, po.exf, extra)
    if n == 0:
        proj_exf = max(TODAY, po.exf) if po.exf >= TODAY else TODAY
    else:
        proj_exf = next_wd(nth_wd(TODAY, n, extra))
    bott = None
    for s in stages:
        if info[s]["binding"] == "own":
            bott = s
    req = None
    if bott:
        later = stages[stages.index(bott) + 1:]
        denom = avail - sum(STAGE_LAG[s] for s in later)
        req = info[bott]["rem"] / denom if denom > 0 else None
    return {"state": "In production", "finish": finish, "n": n, "avail": avail, "slack_wd": avail - n,
            "proj_exf": proj_exf, "slip": (proj_exf - po.exf).days, "bottleneck": bott,
            "bott_rate": info[bott]["eff"] if bott else None, "req_rate": req, "stages": info}


def project_pre(po: PO, ta: dict) -> dict:
    """Projection for a PO whose knitting has not started: the plan shifts by the start delay."""
    yarn = ta["Yarn in-house"]
    yarn_date = yarn["actual"] or max(yarn["revised"] or yarn["planned"], TODAY)
    knit_planned = ta["Knitting start"]["planned"]
    if ta["Knitting start"]["actual"]:          # started, but not yet in a received report
        proj_knit = ta["Knitting start"]["actual"]
    else:
        proj_knit = max(knit_planned, yarn_date + timedelta(days=YARN_TO_KNIT_DAYS))
        if knit_planned < TODAY:
            proj_knit = max(proj_knit, TODAY)
    delay = max(0, (proj_knit - knit_planned).days)
    proj_exf = po.exf + timedelta(days=delay)
    return {"state": "Pre-production", "finish": None, "n": None, "avail": wd_between(TODAY, po.exf),
            "slack_wd": None, "proj_exf": proj_exf, "slip": delay, "bottleneck": None, "bott_rate": None,
            "req_rate": None, "proj_knit": proj_knit, "stages": {}}


def reported_rows(po: PO) -> list[dict]:
    lr = FAC[po.factory].last_report
    return [x for x in po.rows if x["date"] <= lr]


# --------------------------------------------------------- solve production -
def solve_rows(po: PO, target_n: int, mult: dict, link_delay: int, seed: int,
               n_range: tuple[int, int] | None = None) -> list[dict]:
    """Find the base rate whose reported data projects completion on working day `target_n`
    (or anywhere inside `n_range` when the exact day is not reachable)."""
    n_lo, n_hi = n_range or (target_n, target_n)
    lr = FAC[po.factory].last_report

    def f_of(base):
        rows = simulate(po, base, po.knit_start, LAST_EXPECTED_REPORT, mult, link_delay, seed)
        rep = [x for x in rows if x["date"] <= lr]
        if not rep:
            return rows, math.inf
        return rows, project(po, rep, lr)["finish"]

    lo, hi = 25.0, max(400.0, po.qty / 2.5)
    if target_n <= 0:
        for _ in range(60):
            mid = (lo + hi) / 2
            _, f = f_of(mid)
            lo, hi = (lo, mid) if f <= 0 else (mid, hi)
        return f_of(hi * 1.03)[0]
    goal = target_n - 0.5
    for _ in range(60):
        mid = (lo + hi) / 2
        _, f = f_of(mid)
        lo, hi = (mid, hi) if f > goal else (lo, mid)
    mid = (lo + hi) / 2
    rows, f = f_of(mid)
    ok = lambda f: math.isfinite(f) and (0 if f <= 0 else math.ceil(f - 1e-9)) in range(n_lo, n_hi + 1)  # noqa: E731
    if not ok(f):  # step discontinuity: scan nearby rates
        for k in range(1, 600):
            for b in (mid * (1 + k / 1000), mid * (1 - k / 1000)):
                rows2, f2 = f_of(b)
                if ok(f2):
                    return rows2
        SOLVE_MISSES.append((po.po_no, po.scenario, target_n, n_lo, n_hi, round(f, 2)))
    return rows


SOLVE_MISSES: list = []


def solve_closed(po: PO, done_by: date, mult: dict, seed: int) -> list[dict]:
    """Rows for a shipped PO whose packing finishes on `done_by`."""
    lo, hi = 25.0, po.qty / 2.0

    def done_date(base):
        rows = simulate(po, base, po.knit_start, done_by, mult, 3, seed, until_packed=True)
        return rows, (rows[-1]["date"] if rows and rows[-1]["cum"]["packing"] >= po.qty else None)

    for _ in range(50):
        mid = (lo + hi) / 2
        _, d = done_date(mid)
        lo, hi = (lo, mid) if d is not None else (mid, hi)
    return done_date(hi)[0]


def profile(r: random.Random, link_bottleneck: bool | None = None) -> dict:
    if link_bottleneck is None:
        link_bottleneck = r.random() < 0.55
    if link_bottleneck:
        return {"knitting": r.uniform(1.35, 1.7), "linking": 1.0, "mending": 1.25, "washing": 1.3, "ironing": 1.3, "packing": 1.3}
    return {"knitting": 1.0, "linking": r.uniform(1.2, 1.4), "mending": 1.3, "washing": 1.35, "ironing": 1.35, "packing": 1.35}


# ---------------------------------------------------------------- generate --
def main():
    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / "demo_uploads").mkdir(exist_ok=True)

    # ---- closed (shipped) history --------------------------------------
    monthly = [((2025, 10), 55), ((2025, 11), 110), ((2025, 12), 80), ((2026, 1), 70), ((2026, 2), 60),
               ((2026, 3), 40), ((2026, 4), 30), ((2026, 5), 25), ((2026, 6), 35), ((2026, 7), 90),
               ((2026, 8), 110), ((2026, 9), 110), ((2026, 10), 60)]
    closed: list[PO] = []
    weights = [f.weight for f in FACTORIES]
    for (y, m), n in monthly:
        a = date(y, m, 15 if (y, m) == (2025, 10) else 1)
        b = date(2026, 10, 14) if (y, m) == (2026, 10) else (date(y + (m == 12), m % 12 + 1, 1) - ONE)
        for _ in range(n):
            exf = rand_date(a, b, rng)
            fac = rng.choices(FACTORIES, weights)[0]
            closed.append(build_po("", exf, fac, rng, open_=False))
    closed.sort(key=lambda p: p.exf)

    # Lateness: exactly 4 late in 1-14 Oct 2026; per-factory OTD targets over 12 months.
    oct_closed = [p for p in closed if p.exf >= date(2026, 10, 1)]
    late_set = set()
    for code in ["GRL", "GRL", "IRB", "SLM"]:
        cand = [p for p in oct_closed if p.factory == code and id(p) not in late_set and p.exf >= date(2026, 10, 3)]
        if not cand:
            cand = [p for p in oct_closed if id(p) not in late_set and p.exf >= date(2026, 10, 3)]
            cand[0].factory = code
        late_set.add(id(cand[0]))
    for f in FACTORIES:
        mine = [p for p in closed if p.factory == f.code]
        n_late = round(len(mine) * (1 - f.otd_target))
        already = sum(1 for p in mine if id(p) in late_set)
        pool = [p for p in mine if id(p) not in late_set and p.exf < date(2026, 10, 1)]
        for p in rng.sample(pool, max(0, min(len(pool), n_late - already))):
            late_set.add(id(p))
    for p in closed:
        p.scenario = "closed"
        if id(p) in late_set:
            big = FAC[p.factory].risky
            slip = rng.choice([1, 2, 2, 3, 3, 4, 5, 6, 7, 9, 12] if big else [1, 1, 2, 2, 3, 4, 5, 6])
            p.actual_exf = not_friday(p.exf + timedelta(days=slip))
            if p.actual_exf <= p.exf:
                p.actual_exf = p.exf + timedelta(days=2) if (p.exf + timedelta(days=1)).weekday() == 4 else p.exf + timedelta(days=1)
            if p.actual_exf > LAST_EXPECTED_REPORT:   # October late shipments must already be out
                p.actual_exf = LAST_EXPECTED_REPORT
                p.exf = min(p.exf, prev_wd(prev_wd(LAST_EXPECTED_REPORT)))
        else:
            p.actual_exf = not_friday(p.exf - timedelta(days=rng.choice([0, 0, 0, 0, 0, 0, 0, 1, 2, 3])))

    # ---- open order book -------------------------------------------------
    def dates_in(a, b, n):
        return sorted(rand_date(a, b, rng) for _ in range(n))

    open_dates = ([date(2026, 9, 24), date(2026, 9, 28)] +
                  [date(2026, 10, 3), date(2026, 10, 6), date(2026, 10, 8), date(2026, 10, 11), date(2026, 10, 13)] +
                  dates_in(date(2026, 10, 15), date(2026, 10, 31), 94) +      # + hero = 95
                  dates_in(date(2026, 11, 1), date(2026, 11, 30), 119) +      # + donor
                  dates_in(date(2026, 12, 1), date(2026, 12, 31), 95) +
                  dates_in(date(2027, 1, 1), date(2027, 1, 31), 60) +
                  dates_in(date(2027, 2, 1), date(2027, 2, 28), 35))
    opn: list[PO] = []
    for exf in open_dates:
        if exf >= date(2026, 12, 21):
            facs = [f for f in FACTORIES if set(f.gauges) & {12, 14}]
        else:
            facs = FACTORIES
        fac = rng.choices(facs, [f.weight for f in facs])[0]
        opn.append(build_po("", exf, fac, rng, open_=True))

    hero_product = ([p for p in PRODUCTS if p[1] == "V-neck Cardigan"][0], 7, YARNS_AW[0])
    hero = build_po("", date(2026, 10, 29), FAC["GRL"], rng, True, hero_product, qty=9600, fob=16.50)
    hero.dept, hero.weight, hero.scenario = "Women", 0.72, "hero"
    donor_product = ([p for p in PRODUCTS if p[0] == "Men" and p[1] == "Crew-neck Pullover"][0], 7, YARNS_AW[2])
    donor = build_po("", date(2026, 11, 10), FAC["GRL"], rng, True, donor_product, qty=14400, fob=11.20)
    donor.scenario = "donor"
    opn += [hero, donor]

    # Fixed risk scenarios
    late_pos = [p for p in opn if p.exf < TODAY]
    for p, code in zip(late_pos, ["GRL", "IRB", "GRL", "SLM", "IRB", None, None]):
        p.scenario = "late"
        if code:
            p.factory = code
            p.merch = FAC[code].merch
    october = [p for p in opn if TODAY <= p.exf <= date(2026, 10, 31) and p.scenario == "ok"]
    rng.shuffle(october)
    plan = ["GRL"] * 3 + ["IRB"] * 3 + ["SLM"] * 3 + [None] * 3   # + hero = 13 predicted late
    for p, code in zip(october[:12], plan):
        p.scenario = "pred_late"
        if code:
            p.factory = code
            p.merch = FAC[code].merch
        elif FAC[p.factory].risky:
            p.factory = "KST"
    for p, code in zip(october[12:16], ["GRL", "IRB", "SLM", "GRL"]):
        p.scenario = "tight"
        p.factory = code
        p.merch = FAC[code].merch
    # Rebuild product specs for POs whose factory changed (gauge compatibility)
    for p in opn:
        if p.scenario in ("hero", "donor"):
            continue
        fac = FAC[p.factory]
        if p.gauge not in fac.gauges:
            new = build_po("", p.exf, fac, rng, True)
            for attr in ("dept", "product", "gauge", "yarn", "yarn_short", "wash", "weight", "knit_min", "link_std", "fob"):
                setattr(p, attr, getattr(new, attr))

    # Scale quantities and prices to the story totals (hero and donor stay fixed)
    fixed = [hero, donor]
    others = [p for p in opn if p not in fixed]
    qf = (TARGET_OPEN_PCS - sum(p.qty for p in fixed)) / sum(p.qty for p in others)
    for p in others:
        p.qty = max(1200, int(round(p.qty * qf / 12)) * 12)
    ff = (TARGET_OPEN_FOB - sum(p.value for p in fixed)) / sum(p.value for p in others)
    for p in others:
        p.fob = round(p.fob * ff, 2)

    # Watch scenario for some on-time POs
    for p in opn:
        if p.scenario == "ok" and rng.random() < 0.12:
            p.scenario = "watch"

    # Value-at-risk tuning: add Nov/Dec POs to the predicted-late set until both pools hit target
    def var_of(codes_in):
        return sum(p.value for p in opn if p.scenario in ("late", "pred_late", "tight", "hero")
                   and (p.factory in RISKY) == codes_in)
    targets = {True: TARGET_VAR * TARGET_VAR_RISKY_SHARE, False: TARGET_VAR * (1 - TARGET_VAR_RISKY_SHARE)}
    for risky in (True, False):
        cands = [p for p in opn if p.scenario in ("ok", "watch") and (p.factory in RISKY) == risky
                 and date(2026, 11, 1) <= p.exf <= date(2026, 12, 31) and p is not donor]
        rng.shuffle(cands)
        if not risky:   # spread over factories
            by_f = defaultdict(list)
            for p in cands:
                by_f[p.factory].append(p)
            cands = [p for grp in zip(*[v for v in by_f.values() if len(v) >= 3]) for p in grp]
        for p in cands:
            gap = targets[risky] - var_of(risky)
            if gap < 12_000:
                break
            if p.value <= gap + 6_000:
                p.scenario = "pred_late"
        gap = targets[risky] - var_of(risky)
        if abs(gap) >= 1:   # trim the last few dollars with one PO's price
            p = max((x for x in opn if x.scenario == "pred_late" and (x.factory in RISKY) == risky
                     and x.exf.month in (11, 12)), key=lambda x: x.qty)
            p.fob = round(p.fob + gap / p.qty, 2)

    # PO numbers by PO date, styles and colours
    everything = closed + opn
    everything.sort(key=lambda p: (p.po_date, p.exf))
    n = 71_004_000
    for p in everything:
        n += rng.randint(3, 9)
        p.po_no = str(n)
    assign_styles(everything, rng)

    # ---- T&A, production simulation ---------------------------------------
    ta_rows, dpr_rows, demo_rows = [], [], []
    ta_by_po = {}
    for p in everything:
        r = random.Random(f"{SEED}-{p.po_no}")
        sc = p.scenario
        if sc == "hero":
            p.yarn_delay = 8
        elif sc == "tight":
            p.yarn_delay = r.randint(8, 10)
        elif sc in ("pred_late", "late"):
            p.yarn_delay = r.randint(4, 9)
        else:   # within the plan's buffer between yarn in-house and knitting start
            buffer = (p.planned_knit - p.planned_yarn).days - YARN_TO_KNIT_DAYS
            p.yarn_delay = min(buffer, r.choice([-2, -1, 0, 0, 0, 0, 1, 1, 2, 3]))
        if sc in ("pred_late", "late", "tight") and r.random() < 0.35:
            p.pp_delay = r.randint(4, 6)
        else:
            p.pp_delay = r.choice([-1, 0, 0, 0, 1, 2])
        yarn_actual = not_friday(p.planned_yarn + timedelta(days=p.yarn_delay))
        p.knit_start = max(p.planned_knit, yarn_actual + timedelta(days=YARN_TO_KNIT_DAYS))
        if sc not in ("hero",) and r.random() < 0.4:
            p.knit_start += ONE
        if p.knit_start.weekday() == 4:
            p.knit_start += ONE
        if p.open and sc in ("ok", "watch") and p.planned_knit < TODAY <= p.knit_start:
            p.knit_start = prev_wd(TODAY)        # on-plan POs have started by now
        if sc == "hero":
            p.knit_start = date(2026, 9, 24)

        # Pre-production predicted-late via a delayed yarn delivery
        pre = p.open and p.knit_start > LAST_EXPECTED_REPORT
        if p.open and sc == "pred_late" and pre:
            p.yarn_revised = not_friday(p.planned_yarn + timedelta(days=r.randint(6, 9)))

        # Simulate
        if not p.open and p.knit_start <= LAST_EXPECTED_REPORT and prev_wd(p.actual_exf) >= DPR_START:
            p.rows = solve_closed(p, prev_wd(p.actual_exf), profile(r), seed=r.randint(0, 10**9))
        elif p.open and not pre:
            seed = r.randint(0, 10**9)
            avail = wd_between(TODAY, p.exf)
            if sc == "hero":
                link_days = wd_list(date(2026, 10, 7), date(2026, 10, 14))
                forced = dict(zip(link_days, [330, 390, 415, 420, 410, 420, 415]))
                mult = {"knitting": 1.0, "linking": 1.0, "mending": 1.3, "washing": 1.3, "ironing": 1.3, "packing": 1.3}
                p.rows = simulate(p, 520, p.knit_start, LAST_EXPECTED_REPORT, mult, 3, seed, forced_link=forced)
                for x in p.rows:
                    if x["out"]["linking"]:
                        x["lmc"] = {330: 10, 390: 12}.get(x["out"]["linking"], 13)
                    x["kmc"] = 26 if x["out"]["knitting"] else 0
            else:
                if sc == "watch" and avail < 4:
                    p.scenario = sc = "ok"
                if sc in ("ok", "donor"):
                    tn, rg = max(0, avail - (8 if sc == "donor" else r.randint(6, 16))), (0, max(0, avail - 6))
                elif sc == "watch":
                    tn, rg = avail - r.randint(3, 5), (max(1, avail - 5), avail - 3)
                elif sc == "tight":
                    tn, rg = avail - r.randint(1, 2), (max(1, avail - 2), avail - 1)
                elif sc == "pred_late":
                    tn, rg = avail + r.choice([1, 2, 2, 3, 4, 5, 6, 7]), (avail + 1, avail + 7)
                else:  # late (past ex-factory)
                    tn, rg = r.randint(2, 6), (1, 7)
                mult = profile(r, link_bottleneck=True if sc == "donor" else None)
                if p.knit_start > FAC[p.factory].last_report:   # started, nothing reported yet
                    p.rows = simulate(p, max(60.0, p.qty / max(avail, 10)), p.knit_start, LAST_EXPECTED_REPORT, mult, 3, seed)
                else:
                    p.rows = solve_rows(p, tn, mult, 3, seed, rg)
        if p.rows:
            k = [x["date"] for x in p.rows if x["out"]["knitting"] > 0]
            l_ = [x["date"] for x in p.rows if x["out"]["linking"] > 0]
            p.knit_start = k[0] if k else p.knit_start
            p.link_start = l_[0] if l_ else None

        # T&A rows
        ta = {}
        for seq, ms, before, resp in TA_TEMPLATE:
            planned = p.exf if ms == "Ex-factory" else (prev_wd(p.exf) if ms == "Final inspection" else not_friday(p.exf - timedelta(days=before)))
            actual = revised = None
            remark = ""
            if ms == "Yarn in-house":
                if p.yarn_revised:
                    revised = p.yarn_revised
                    remark = r.choice(["Dyed yarn delayed at spinner (shade re-dye)", "Yarn shipment rolled to next vessel",
                                       "Supplier capacity issue - partial lot only"])
                else:
                    actual = yarn_actual
                    if p.yarn_delay >= 4:
                        remark = r.choice(["Dyed yarn delayed at spinner (shade re-dye)", "Yarn held at customs 4 days",
                                           "Lab dip re-approval needed for one colour"])
            elif ms == "Knitting start":
                actual = p.knit_start if (p.rows or not p.open) else None
            elif ms == "Linking start":
                actual = p.link_start if p.rows else (not_friday(p.knit_start + timedelta(days=5)) if not p.open else None)
                if sc == "hero":
                    remark = "Linking machines tied up on other POs; 13 M/C allocated from 7 Oct"
            elif ms == "PP sample approval":
                actual = not_friday(planned + timedelta(days=p.pp_delay))
                if p.pp_delay >= 4:
                    remark = "PP sample rejected once (measurement); resubmitted"
            elif ms in ("Final inspection", "Ex-factory"):
                if not p.open:
                    actual = prev_wd(p.actual_exf) if ms == "Final inspection" else p.actual_exf
            else:
                actual = not_friday(planned + timedelta(days=r.choice([-2, -1, 0, 0, 0, 1, 2])))
            if actual and actual > LAST_EXPECTED_REPORT:
                actual = None
            if actual:
                status = "Done late" if actual > planned else "Done"
            else:
                status = "Overdue" if planned < TODAY else "Pending"
            ta[ms] = {"planned": planned, "revised": revised, "actual": actual}
            ta_rows.append({"PO No": p.po_no, "Factory Code": p.factory, "Seq": seq, "Milestone": ms,
                            "Responsible": resp, "Planned Date": planned, "Revised Date": revised,
                            "Actual Date": actual, "Status": status, "Remarks": remark})
        ta_by_po[p.po_no] = ta

        # DPR rows (only what the factory reported, from DPR_START)
        for x in p.rows:
            if x["date"] < DPR_START:
                continue
            row = dpr_row(p, x)
            if x["date"] > FAC[p.factory].last_report:     # not reported yet
                if p.factory in ("PBB", "HMR") and x["date"] == LAST_EXPECTED_REPORT:
                    demo_rows.append((p.factory, row))
                continue
            dpr_rows.append(row)

    # ---- shipments and inspections ----------------------------------------
    ship_rows, insp_rows = [], []
    insp_n = 0
    for p in sorted(closed, key=lambda x: x.actual_exf):
        r = random.Random(f"ship-{p.po_no}")
        slip = (p.actual_exf - p.exf).days
        short = 0 if r.random() < 0.8 else r.uniform(0.005, 0.02)
        shipped = int(p.qty * (1 - short))
        cartons = math.ceil(shipped / (PCS_PER_CARTON[p.gauge] * (2 if p.dept == "Kids" else 1)))
        gw = round(shipped * p.weight + cartons * 1.2, 1)
        if slip >= 5 and r.random() < 0.5:
            mode = "Air"
        elif r.random() < 0.04:
            mode = "Sea-Air"
        else:
            mode = "Sea"
        dest = next(d for d in DESTINATIONS if d[0] == p.destination)
        if mode == "Air":
            etd, eta = p.actual_exf + timedelta(days=r.randint(1, 2)), None
            eta = etd + timedelta(days=r.randint(3, 5))
            vessel, pol = f"Air consolidation AWB {r.randint(1000000, 9999999)}", "Dhaka (HSIA)"
            cost = round(gw * AIR_USD_PER_KG, 2)
            bearer = "Factory" if r.random() < 0.8 else "Karbar"
        else:
            etd = p.actual_exf + timedelta(days=r.randint(4, 7))
            eta = etd + timedelta(days=r.randint(*dest[2]) if mode == "Sea" else r.randint(18, 22))
            vessel = f"{r.choice(VESSELS)} V.{etd.strftime('%y%m')}{r.choice('EW')}"
            pol, cost, bearer = "Chattogram", None, "Karbar"
        status = "Delivered" if eta <= TODAY else ("In transit" if etd <= TODAY else "At port")
        ship_rows.append({
            "Shipment ID": f"SH-{p.actual_exf:%y%m}-{len(ship_rows) + 1:04d}", "PO No": p.po_no,
            "Factory Code": p.factory, "Planned Ex-factory": p.exf, "Actual Ex-factory": p.actual_exf,
            "Shipped Qty": shipped, "Cartons": cartons, "Gross Weight kg": gw, "Ship Mode": mode,
            "Port of Loading": pol, "Destination": p.destination, "ETD": etd, "ETA": eta,
            "Vessel / Flight": vessel, "Invoice No": f"KB-INV-{p.actual_exf:%y%m}-{r.randint(1000, 9999)}",
            "Invoice Value USD": round(shipped * p.fob, 2), "Air Freight Cost USD": cost,
            "Freight Cost Borne By": bearer, "Shipment Status": status,
            "Remarks": (f"Shipped {slip} days late" + (" - air freight at factory cost" if mode == "Air" and bearer == "Factory" else "")) if slip > 0 else "",
        })
        # final inspection (with a failed first attempt for weak factories)
        fail_p = 0.18 if p.factory == "IRB" else (0.08 if FAC[p.factory].risky else 0.04)
        fi_date = prev_wd(p.actual_exf)
        if r.random() < fail_p:
            insp_n += 1
            insp_rows.append(insp_row(insp_n, prev_wd(prev_wd(fi_date)), "Final", p, p.qty, False, r))
            insp_n += 1
            insp_rows.append(insp_row(insp_n, fi_date, "Final re-inspection", p, p.qty, True, r))
        else:
            insp_n += 1
            insp_rows.append(insp_row(insp_n, fi_date, "Final", p, p.qty, True, r))
    for p in opn:
        if not p.rows:
            continue
        r = random.Random(f"insp-{p.po_no}")
        rep = reported_rows(p)
        fail_p = 0.30 if p.factory == "IRB" else (0.10 if FAC[p.factory].risky else 0.05)
        for stage, frac, typ in (("linking", 0.15, "Inline"), ("packing", 0.40, "Pre-final")):
            hit = next((x for x in rep if x["cum"][stage] >= frac * p.qty), None)
            if hit and next_wd(hit["date"]) <= LAST_EXPECTED_REPORT:
                ok = (p.scenario == "hero") or r.random() >= fail_p
                insp_n += 1
                insp_rows.append(insp_row(insp_n, next_wd(hit["date"]), typ, p, hit["cum"][stage], ok, r))
    insp_rows.sort(key=lambda x: x["Inspection Date"])
    for i, x in enumerate(insp_rows, 1):
        x["Inspection ID"] = f"QA-{x['Inspection Date']:%y%m}-{i:05d}"

    # ---- factory master capacity from simulated load -------------------------
    fac_rows = factory_master(everything)

    # ---- predictions (answer key) -----------------------------------------
    ship_df = pd.DataFrame(ship_rows)
    insp_df = pd.DataFrame(insp_rows)
    preds = predict_all(opn, ta_by_po, ship_df, insp_df)
    report = summarise(opn, preds, closed, ship_df)
    whatif = hero_whatif(hero, donor)

    # ---- write files ------------------------------------------------------
    write_order_book(everything)
    write_book(OUT / "02_TA_Calendar.xlsx", [("T&A Calendar", pd.DataFrame(ta_rows), TA_COLS)], DATE_COLS_TA)
    dpr_df = pd.DataFrame(dpr_rows).sort_values(["Report Date", "Factory Code", "PO No"])
    write_book(OUT / "03_Factory_Daily_Production_Report.xlsx", [("Daily Production", dpr_df, DPR_COLS)], ["Report Date"])
    write_book(OUT / "04_Inspection_Log.xlsx", [("Inspections", insp_df[list(INSP_COLS)], INSP_COLS)], ["Inspection Date"])
    write_book(OUT / "05_Factory_Master.xlsx", [("Factories", pd.DataFrame(fac_rows), FAC_COLS)], ["Supplier Since", "Last Social Audit"])
    write_book(OUT / "06_Shipment_Log.xlsx", [("Shipments", ship_df, SHIP_COLS)],
               ["Planned Ex-factory", "Actual Ex-factory", "ETD", "ETA"])
    write_demo_uploads(demo_rows)
    write_answer_key(report, preds, whatif)
    print_summary(report, whatif, len(dpr_rows), len(ta_rows), len(insp_rows), len(ship_rows), len(everything))


def dpr_row(p: PO, x: dict) -> dict:
    o, c = x["out"], x["cum"]
    row = {"Report Date": x["date"], "Factory Code": p.factory, "Factory Name": FAC[p.factory].name,
           "PO No": p.po_no, "Style No": p.style_no, "Order Qty": p.qty,
           "Knitting Day": o["knitting"], "Knitting Cum": c["knitting"], "Knitting M/C": x["kmc"],
           "Linking Day": o["linking"], "Linking Cum": c["linking"], "Linking M/C": x["lmc"],
           "Trimming & Mending Day": o["mending"], "Trimming & Mending Cum": c["mending"],
           "Washing Day": o.get("washing") if p.wash == "Y" else None,
           "Washing Cum": c.get("washing") if p.wash == "Y" else None,
           "Ironing Day": o["ironing"], "Ironing Cum": c["ironing"],
           "Packing Day": o["packing"], "Packing Cum": c["packing"], "Remarks": ""}
    r = random.Random(f"rmk-{p.po_no}-{x['date']}")
    if p.scenario == "hero" and x["date"] == date(2026, 10, 7):
        row["Remarks"] = "Linking started late - linking M/C busy on other POs"
    elif p.scenario == "hero" and x["date"] == date(2026, 10, 13):
        row["Remarks"] = "Need 6 more linking M/C to recover; requested from planning"
    elif p.scenario in ("pred_late", "late", "tight") and r.random() < 0.12:
        row["Remarks"] = r.choice(["Linking operators short (absent)", "Power outage 2 hrs - generator on",
                                   "Waiting for trims (buttons)", "Yarn shortage in one colour",
                                   "Machines shifted to urgent PO", "Washing re-run for hand-feel"])
    elif r.random() < 0.015:
        row["Remarks"] = r.choice(["Running as plan", "Overtime 2 hrs", "Inline QC done"])
    return row


def aql(lot: int):
    for hi, sample, ac_maj, ac_min in ((500, 50, 3, 5), (1200, 80, 5, 7), (3200, 125, 7, 10),
                                       (10000, 200, 10, 14), (35000, 315, 14, 21)):
        if lot <= hi:
            return sample, ac_maj, ac_min
    return 500, 21, 21


def insp_row(n, d, typ, p: PO, lot, ok, r):
    sample, ac, ac_min = aql(lot)
    major = r.randint(0, ac) if ok else r.randint(ac + 1, ac + 6)
    spec = round(p.weight * 1000 - 40)
    return {"Inspection ID": "", "Inspection Date": d, "Inspection Type": typ, "PO No": p.po_no,
            "Factory Code": p.factory, "Inspector": r.choice(INSPECTORS), "Lot Qty": int(lot), "AQL Level": "II / 2.5 / 4.0",
            "Sample Size": sample, "Major Accept (Ac)": ac, "Critical Found": 0, "Major Found": major,
            "Minor Found": r.randint(0, ac_min), "Result": "Pass" if ok else "Fail",
            "Main Defect": r.choice(DEFECTS) if major else "",
            "Measurement Check": "Pass" if ok or r.random() < 0.5 else "Fail",
            "Spec Weight g": spec, "Avg Weight g": spec + r.randint(-8, 10) if ok else spec - r.randint(10, 25),
            "Remarks": "" if ok else "Re-inspection required after 100% check"}


def factory_master(pos: list[PO]) -> list[dict]:
    use_k = defaultdict(lambda: defaultdict(list))
    use_l = defaultdict(list)
    window = wd_list(date(2026, 10, 1), LAST_EXPECTED_REPORT)
    for d in window:
        day_k = defaultdict(lambda: defaultdict(int))
        day_l = defaultdict(int)
        for p in pos:
            for x in p.rows:
                if x["date"] == d:
                    day_k[p.factory][p.gauge] += x["kmc"]
                    day_l[p.factory] += x["lmc"]
        for f in FACTORIES:
            for g in f.gauges:
                use_k[f.code][g].append(day_k[f.code][g])
            use_l[f.code].append(day_l[f.code])
    rows = []
    r = random.Random(SEED + 7)
    for f in FACTORIES:
        util = 1.0 if f.risky else r.uniform(0.78, 0.88)
        mc = {}
        for g in (3, 5, 7, 12, 14):
            if g in f.gauges:
                peak = max(use_k[f.code][g]) if use_k[f.code][g] else 0
                mc[g] = max(40, int(math.ceil(peak / util / 10.0)) * 10)
            else:
                mc[g] = 0
        lpeak = max(use_l[f.code]) if use_l[f.code] else 0
        rows.append({
            "Factory Code": f.code, "Factory Name": f.name, "Area": f.area, "District": f.district,
            "Address": f"Plot {r.randint(2, 98)}, {f.area}, {f.district}", "Supplier Since": date(r.randint(2014, 2023), r.randint(1, 12), 1),
            "Gauges": ", ".join(f"{g}GG" for g in f.gauges),
            "Knitting M/C 3GG": mc[3], "Knitting M/C 5GG": mc[5], "Knitting M/C 7GG": mc[7],
            "Knitting M/C 12GG": mc[12], "Knitting M/C 14GG": mc[14], "Total Knitting M/C": sum(mc.values()),
            "Linking M/C": max(60, int(math.ceil(lpeak / util / 10.0)) * 10),
            "Washing": "In-house" if r.random() < 0.6 else "Subcontract",
            "Knitting Shifts": 2, "Knitting Hours/Day": 20, "Linking Hours/Day": 10,
            "Workforce": r.randint(900, 3200) if f.weight >= 4 else r.randint(500, 1400),
            "Certifications": ", ".join(r.sample(["amfori BSCI", "WRAP", "OEKO-TEX STeP", "GOTS", "ISO 14001", "Higg FEM verified"], r.randint(2, 4))),
            "amfori BSCI Rating": r.choice(["A", "B", "B", "C"]) if not f.risky else r.choice(["B", "C"]),
            "Last Social Audit": date(2026, r.randint(1, 9), r.randint(1, 28)),
            "Factory Contact": r.choice(["Md. Rafiqul Islam", "Shahidul Alam", "Nazmul Haque", "Sharmin Akter",
                                         "Abdullah Al Mamun", "Kazi Fahim", "Rezaul Karim", "Tahmina Begum"]),
            "Contact Title": r.choice(["GM - Merchandising", "Head of Production", "Director", "AGM - Commercial"]),
            "Karbar Merchandiser": f.merch, "Status": "Active",
        })
    return rows


# ---------------------------------------------------------------- answer key
def predict_all(opn, ta_by_po, ship_df, insp_df) -> list[dict]:
    hist = ship_df[(ship_df["Actual Ex-factory"] >= HISTORY_START) & (ship_df["Actual Ex-factory"] <= LAST_EXPECTED_REPORT)].copy()
    hist["slip"] = [(a - b).days for a, b in zip(hist["Actual Ex-factory"], hist["Planned Ex-factory"])]
    otd = {f.code: float((hist[hist["Factory Code"] == f.code]["slip"] <= 0).mean()) for f in FACTORIES}
    slips = {f.code: list(hist[hist["Factory Code"] == f.code]["slip"]) for f in FACTORIES}
    q90 = insp_df[insp_df["Inspection Date"] >= TODAY - timedelta(days=90)]
    pass90 = {f.code: float((q90[q90["Factory Code"] == f.code]["Result"] == "Pass").mean()) if (q90["Factory Code"] == f.code).any() else 1.0
              for f in FACTORIES}
    last_insp = insp_df.sort_values("Inspection Date").groupby("PO No").last()["Result"].to_dict()
    out = []
    for p in opn:
        ta = ta_by_po[p.po_no]
        rep = reported_rows(p)
        pr = project(p, rep, FAC[p.factory].last_report) if rep and rep[-1]["cum"]["knitting"] > 0 else project_pre(p, ta)
        slip = pr["slip"]
        # risk score
        if pr["n"] == 0:
            sched = 0                       # fully packed, waiting for ex-factory
        elif slip > 0:
            sched = 35 + min(15, 1.5 * slip)
        elif pr["state"] == "In production" and pr["slack_wd"] <= 2:
            sched = 25
        elif pr["state"] == "In production" and pr["slack_wd"] <= 5:
            sched = 12
        else:
            sched = 0
        y = ta["Yarn in-house"]
        y_late = ((y["actual"] or max(y["revised"] or y["planned"], TODAY)) - y["planned"]).days
        ppa = ta["PP sample approval"]
        pp_late = ((ppa["actual"] or TODAY) - ppa["planned"]).days if (ppa["actual"] or ppa["planned"] < TODAY) else 0
        tna = (12 if y_late >= 8 else 8 if y_late >= 4 else 4 if y_late >= 1 else 0) + (8 if pp_late >= 3 else 0)
        tna = min(20, tna)
        rel = min(15.0, (1 - otd[p.factory]) * 50)
        qual = 10 if last_insp.get(p.po_no) == "Fail" else (5 if pass90[p.factory] < 0.85 else 0)
        stale = wd_between(next_wd(FAC[p.factory].last_report), next_wd(LAST_EXPECTED_REPORT)) if pr["state"] == "In production" else 0
        fresh = 0 if stale == 0 else (3 if stale == 1 else 5)
        score = round(min(100, sched + tna + rel + qual + fresh))
        band = "On track" if score < 25 else "Watch" if score < 50 else "At risk" if score < 75 else "Critical"
        if slip > 0 and band in ("On track", "Watch"):
            band = "At risk"
        if slip >= 7:
            band = "Critical"
        if p.exf < TODAY:
            band = "Late"
        s_cal = -slip
        hs = slips[p.factory]
        prob = min(0.98, max(0.02, sum(1 for v in hs if v <= s_cal) / len(hs))) if hs else 0.5
        at_risk = band in ("At risk", "Critical", "Late")
        late_flag = slip > 0 or p.exf < TODAY
        drivers = []
        if pr.get("bottleneck") and pr.get("req_rate") and pr["bott_rate"] < pr["req_rate"]:
            drivers.append(f"{STAGE_LABEL[pr['bottleneck']]} {pr['bott_rate']:,.0f} pcs/day vs {pr['req_rate']:,.0f} needed")
        if y["revised"] and not y["actual"]:
            drivers.append(f"Yarn in-house revised to {y['revised']:%d %b} ({y_late} days late)")
        elif y_late >= 4:
            drivers.append(f"Yarn in-house {y_late} days late")
        if pp_late >= 3:
            drivers.append(f"PP sample approval {pp_late} days late")
        if otd[p.factory] < 0.8:
            drivers.append(f"Factory OTD {otd[p.factory]:.0%} (12 months)")
        if qual == 10:
            drivers.append("Last inspection failed")
        if fresh:
            drivers.append(f"No report since {FAC[p.factory].last_report:%d %b}")
        out.append({
            "PO No": p.po_no, "Factory Code": p.factory, "Factory Name": FAC[p.factory].name, "Department": p.dept,
            "Style No": p.style_no, "Style Name": p.style_name, "Gauge": p.gauge, "Order Qty": p.qty,
            "FOB Value USD": p.value, "Planned Ex-factory": p.exf, "State": pr["state"],
            "Bottleneck Stage": STAGE_LABEL.get(pr["bottleneck"], "") if pr["bottleneck"] else "",
            "Bottleneck Rate pcs/day": round(pr["bott_rate"], 1) if pr["bott_rate"] else None,
            "Required Rate pcs/day": round(pr["req_rate"], 1) if pr["req_rate"] else None,
            "Projected Finish (WD)": round(pr["finish"], 2) if pr["finish"] is not None else None,
            "Completion Day N": pr["n"], "Available WD": pr["avail"], "Slack WD": pr["slack_wd"],
            "Projected Ex-factory": pr["proj_exf"], "Slip Days": slip,
            "Score Schedule": round(sched, 1), "Score T&A": tna, "Score Factory OTD": round(rel, 1),
            "Score Quality": qual, "Score Freshness": fresh, "Risk Score": score, "Band": band,
            "On-time Probability": round(prob, 2), "Value at Risk USD": p.value if at_risk else 0.0,
            "Air-freight Exposure USD": round(p.qty * p.weight * (AIR_USD_PER_KG - SEA_USD_PER_KG), 2) if late_flag else 0.0,
            "Drivers": "; ".join(drivers), "Merchandiser": p.merch, "Seed Scenario": p.scenario,
            "_otd": otd[p.factory], "_pass90": pass90[p.factory],
        })
    return out


def summarise(opn, preds, closed, ship_df) -> dict:
    df = pd.DataFrame(preds)
    octo = df[(df["Planned Ex-factory"] >= date(2026, 10, 1)) & (df["Planned Ex-factory"] <= date(2026, 10, 31))]
    oct_ship = ship_df[(ship_df["Planned Ex-factory"] >= date(2026, 10, 1)) & (ship_df["Planned Ex-factory"] <= date(2026, 10, 31))]
    oct_ship_on = int(((oct_ship["Actual Ex-factory"] - oct_ship["Planned Ex-factory"]).apply(lambda x: x.days) <= 0).sum())
    oct_open_on = int(((octo["Slip Days"] <= 0) & (octo["Band"] != "Late")).sum())
    oct_total = len(octo) + len(oct_ship)
    var = df["Value at Risk USD"].sum()
    by_f = df.groupby("Factory Code")["Value at Risk USD"].sum().sort_values(ascending=False)
    top3 = by_f.head(3)
    reported = sum(1 for f in FACTORIES if f.last_report == LAST_EXPECTED_REPORT)
    weeks = []
    w0 = date(2026, 10, 10)  # Saturday; the Bangladesh work week runs Saturday to Thursday
    for i in range(8):
        a, b = w0 + timedelta(days=7 * i), w0 + timedelta(days=7 * i + 6)
        wk = df[(df["Planned Ex-factory"] >= a) & (df["Planned Ex-factory"] <= b)]
        row = {"Week": f"{a:%d %b} - {b:%d %b}", "POs": len(wk), "Pcs": int(wk["Order Qty"].sum()), "FOB USD": round(wk["FOB Value USD"].sum(), 2)}
        for band in ("On track", "Watch", "At risk", "Critical", "Late"):
            row[f"{band} USD"] = round(wk[wk["Band"] == band]["FOB Value USD"].sum(), 2)
        weeks.append(row)
    fac = []
    for f in FACTORIES:
        sub = df[df["Factory Code"] == f.code]
        fac.append({"Factory Code": f.code, "Factory Name": f.name, "Open POs": len(sub),
                    "Open FOB USD": round(sub["FOB Value USD"].sum(), 2),
                    "At risk + Critical + Late POs": int(sub["Band"].isin(["At risk", "Critical", "Late"]).sum()),
                    "Value at Risk USD": round(sub["Value at Risk USD"].sum(), 2),
                    "Share of Value at Risk": round(sub["Value at Risk USD"].sum() / var, 3) if var else 0,
                    "OTD 12m": round(sub["_otd"].iloc[0], 3) if len(sub) else None,
                    "AQL Pass 90d": round(sub["_pass90"].iloc[0], 3) if len(sub) else None,
                    "Last Report": f.last_report, "Reported Today": "Yes" if f.last_report == LAST_EXPECTED_REPORT else "No"})
    return {
        "df": df, "weeks": weeks, "factories": fac,
        "kpis": [
            ("As of", f"{TODAY:%d %b %Y} (reports up to {LAST_EXPECTED_REPORT:%d %b})"),
            ("Open POs", len(df)),
            ("Open qty (pcs)", int(df["Order Qty"].sum())),
            ("Open FOB value (USD)", round(df["FOB Value USD"].sum(), 2)),
            ("October POs (planned ex-factory in Oct)", oct_total),
            ("October shipped so far", len(oct_ship)),
            ("October shipped on time", oct_ship_on),
            ("October open, predicted on time", oct_open_on),
            ("October predicted on-time %", round((oct_ship_on + oct_open_on) / oct_total, 4)),
            ("October open, predicted late (excl. already late)", int(((octo["Slip Days"] > 0) & (octo["Band"] != "Late")).sum())),
            ("Already late (past ex-factory, not shipped)", int((df["Band"] == "Late").sum())),
            ("Value at risk (USD) = FOB of At risk + Critical + Late", round(var, 2)),
            ("POs at risk (At risk + Critical + Late)", int(df["Band"].isin(["At risk", "Critical", "Late"]).sum())),
            ("Top 3 factories by value at risk", ", ".join(f"{k} {v:,.0f}" for k, v in top3.items())),
            ("Top 3 share of value at risk", round(top3.sum() / var, 4)),
            ("Air-freight exposure (USD)", round(df["Air-freight Exposure USD"].sum(), 2)),
            ("Factories reported today", f"{reported} of {len(FACTORIES)}"),
            ("Band counts", ", ".join(f"{k}: {v}" for k, v in df["Band"].value_counts().items())),
            ("Shipping rest of October (POs / pcs / USD)",
             f"{len(octo[octo['Planned Ex-factory'] >= TODAY])} / {int(octo[octo['Planned Ex-factory'] >= TODAY]['Order Qty'].sum()):,} / "
             f"{octo[octo['Planned Ex-factory'] >= TODAY]['FOB Value USD'].sum():,.0f}"),
        ],
    }


def hero_whatif(hero: PO, donor: PO) -> list[dict]:
    lr = FAC[hero.factory].last_report
    rep = reported_rows(hero)
    fridays = frozenset({date(2026, 10, 16), date(2026, 10, 23)})
    mc = 13
    res = []
    for label, mult, extra in [("Current plan", 1.0, frozenset()),
                               ("+6 linking M/C (13 -> 19)", (mc + 6) / mc, frozenset()),
                               ("2 Friday overtime days (16, 23 Oct)", 1.0, fridays),
                               ("+6 linking M/C and 2 Friday overtime days", (mc + 6) / mc, fridays)]:
        pr = project(hero, rep, lr, extra=extra, link_mult=mult)
        res.append({"PO No": hero.po_no, "Option": label, "Linking pcs/day": round(pr["stages"]["linking"]["eff"], 1),
                    "Required pcs/day": round(pr["req_rate"], 1) if pr["req_rate"] else None,
                    "Available WD": pr["avail"], "Completion Day N": pr["n"],
                    "Projected Ex-factory": pr["proj_exf"], "Slip Days": pr["slip"],
                    "On time": "Yes" if pr["slip"] <= 0 else "No",
                    "Air freight avoided USD": round(hero.qty * hero.weight * (AIR_USD_PER_KG - SEA_USD_PER_KG), 2) if pr["slip"] <= 0 else 0})
    drep = reported_rows(donor)
    dmc = drep[-1]["lmc"]
    for label, mult in [("Donor today", 1.0), (f"Donor after giving 6 linking M/C ({dmc} -> {dmc - 6})", (dmc - 6) / dmc)]:
        pr = project(donor, drep, FAC[donor.factory].last_report, link_mult=mult)
        res.append({"PO No": donor.po_no, "Option": label, "Linking pcs/day": round(pr["stages"]["linking"]["eff"], 1),
                    "Required pcs/day": round(pr["req_rate"], 1) if pr["req_rate"] else None,
                    "Available WD": pr["avail"], "Completion Day N": pr["n"],
                    "Projected Ex-factory": pr["proj_exf"], "Slip Days": pr["slip"],
                    "On time": "Yes" if pr["slip"] <= 0 else "No", "Air freight avoided USD": None})
    return res


# ------------------------------------------------------------------ writing -
# Column dictionaries: name -> (type, required, description, example)
PO_COLS = {
    "PO No": ("Text", "Yes", "Karbar purchase order number. Unique.", "71004521"),
    "PO Date": ("Date", "Yes", "Date the PO was issued to the factory.", "12-Jun-2026"),
    "Season": ("Text", "Yes", "Buying season: AW, HO, SS or SU plus year.", "AW26"),
    "Department": ("Text", "Yes", "Men, Women or Kids.", "Women"),
    "Style No": ("Text", "Yes", "Karbar style number.", "KAW26-W-1184"),
    "Style Name": ("Text", "Yes", "Style description.", "Lambswool V-neck Cardigan"),
    "Product Type": ("Text", "Yes", "Garment type.", "V-neck Cardigan"),
    "Gauge": ("Integer", "Yes", "Knitting gauge (GG): 3, 5, 7, 12 or 14.", "7"),
    "Yarn Composition": ("Text", "Yes", "Fibre content.", "100% Lambswool"),
    "Wash Required": ("Y/N", "Yes", "Y if the style goes through washing (softener/enzyme).", "Y"),
    "Weight kg/pc": ("Decimal", "Yes", "Packed garment weight, used for freight cost.", "0.72"),
    "Knitting Minutes/pc": ("Decimal", "Yes", "Standard knitting time for one piece on one machine.", "50.4"),
    "Linking Std pcs/M/C/day": ("Integer", "Yes", "Standard linking output per machine per 10-hour day.", "32"),
    "Factory Code": ("Text", "Yes", "Factory that makes the PO. Must exist in Factory Master.", "GRL"),
    "Factory Name": ("Text", "No", "For readability; Factory Master is the source.", "Greyloom Knitwear Ltd."),
    "Order Qty": ("Integer", "Yes", "Ordered pieces. Equals the sum of the PO Lines.", "9600"),
    "FOB USD/pc": ("Decimal", "Yes", "FOB price per piece in USD.", "16.50"),
    "FOB Value USD": ("Decimal", "Yes", "Order Qty x FOB USD/pc.", "158400.00"),
    "Planned Ex-factory": ("Date", "Yes", "Committed ex-factory date. The on-time target.", "29-Oct-2026"),
    "Planned Ship Mode": ("Text", "Yes", "Sea unless agreed otherwise.", "Sea"),
    "Port of Loading": ("Text", "Yes", "Chattogram for sea.", "Chattogram"),
    "Destination": ("Text", "Yes", "Karbar distribution centre.", "Hamburg, Germany"),
    "Delivery Terms": ("Text", "Yes", "Incoterm.", "FOB Chattogram"),
    "Merchandiser": ("Text", "Yes", "Karbar merchandiser who owns the PO.", "Farhana Rahman"),
}
LINE_COLS = {
    "PO No": ("Text", "Yes", "Links to PO Header.", "71004521"),
    "Colour": ("Text", "Yes", "Colour name.", "Navy"),
    "Colour Code": ("Text", "Yes", "Karbar colour code.", "NAV-410"),
    **{s: ("Integer", "No", f"Pieces in size {s}. Blank when the size is not in the range.", "120") for s in ALL_SIZES},
    "Line Qty": ("Integer", "Yes", "Sum of the size columns.", "3200"),
}
TA_COLS = {
    "PO No": ("Text", "Yes", "Links to the Order Book.", "71004521"),
    "Factory Code": ("Text", "Yes", "Factory that owns the milestone.", "GRL"),
    "Seq": ("Integer", "Yes", "Milestone order, 1 to 11.", "6"),
    "Milestone": ("Text", "Yes", "Yarn booking, Yarn shade approval, Fit sample approval, Size set approval, PP sample approval, "
                  "Yarn in-house, PP meeting, Knitting start, Linking start, Final inspection, Ex-factory.", "Yarn in-house"),
    "Responsible": ("Text", "Yes", "Who must complete it.", "Factory / Yarn supplier"),
    "Planned Date": ("Date", "Yes", "Date from the T&A template (days before ex-factory).", "14-Sep-2026"),
    "Revised Date": ("Date", "No", "New committed date when the factory has warned of a delay.", "20-Oct-2026"),
    "Actual Date": ("Date", "No", "Date it happened. Blank while pending.", "22-Sep-2026"),
    "Status": ("Text", "Yes", "Done, Done late, Pending or Overdue (as of the report date).", "Done late"),
    "Remarks": ("Text", "No", "Reason for a delay.", "Dyed yarn delayed at spinner"),
}
DATE_COLS_TA = ["Planned Date", "Revised Date", "Actual Date"]
DPR_COLS = {
    "Report Date": ("Date", "Yes", "Production date reported.", "14-Oct-2026"),
    "Factory Code": ("Text", "Yes", "Must exist in Factory Master.", "GRL"),
    "Factory Name": ("Text", "No", "For readability.", "Greyloom Knitwear Ltd."),
    "PO No": ("Text", "Yes", "Must exist in the Order Book.", "71004521"),
    "Style No": ("Text", "No", "For readability.", "KAW26-W-1184"),
    "Order Qty": ("Integer", "No", "For readability; the Order Book is the source.", "9600"),
    "Knitting Day": ("Integer", "Yes", "Pieces knitted that day (complete panel sets).", "520"),
    "Knitting Cum": ("Integer", "Yes", "Cumulative knitted. Never decreases.", "8900"),
    "Knitting M/C": ("Integer", "No", "Knitting machines running on the PO that day.", "26"),
    "Linking Day": ("Integer", "Yes", "Pieces linked that day.", "415"),
    "Linking Cum": ("Integer", "Yes", "Cumulative linked. Never above Knitting Cum.", "2800"),
    "Linking M/C": ("Integer", "No", "Linking machines on the PO that day.", "13"),
    "Trimming & Mending Day": ("Integer", "Yes", "Pieces trimmed and mended that day.", "430"),
    "Trimming & Mending Cum": ("Integer", "Yes", "Cumulative. Never above Linking Cum.", "2700"),
    "Washing Day": ("Integer", "No", "Blank when the style has no wash.", "420"),
    "Washing Cum": ("Integer", "No", "Cumulative washed.", "2600"),
    "Ironing Day": ("Integer", "Yes", "Pieces ironed that day.", "410"),
    "Ironing Cum": ("Integer", "Yes", "Cumulative ironed.", "2450"),
    "Packing Day": ("Integer", "Yes", "Pieces packed that day.", "400"),
    "Packing Cum": ("Integer", "Yes", "Cumulative packed. Ready to ship when it equals Order Qty.", "2250"),
    "Remarks": ("Text", "No", "Factory comment.", "Need 6 more linking M/C"),
}
INSP_COLS = {
    "Inspection ID": ("Text", "Yes", "Unique inspection reference.", "QA-2610-01234"),
    "Inspection Date": ("Date", "Yes", "Date of inspection.", "12-Oct-2026"),
    "Inspection Type": ("Text", "Yes", "Inline, Pre-final, Final or Final re-inspection.", "Inline"),
    "PO No": ("Text", "Yes", "Links to the Order Book.", "71004521"),
    "Factory Code": ("Text", "Yes", "Factory inspected.", "GRL"),
    "Inspector": ("Text", "Yes", "Karbar QA inspector.", "Lipi Das"),
    "Lot Qty": ("Integer", "Yes", "Pieces offered for inspection.", "9600"),
    "AQL Level": ("Text", "Yes", "Inspection level / major AQL / minor AQL.", "II / 2.5 / 4.0"),
    "Sample Size": ("Integer", "Yes", "ANSI Z1.4 sample size for the lot.", "200"),
    "Major Accept (Ac)": ("Integer", "Yes", "Maximum major defects to pass.", "10"),
    "Critical Found": ("Integer", "Yes", "Critical defects found. Any critical fails the lot.", "0"),
    "Major Found": ("Integer", "Yes", "Major defects found.", "6"),
    "Minor Found": ("Integer", "Yes", "Minor defects found.", "9"),
    "Result": ("Text", "Yes", "Pass or Fail.", "Pass"),
    "Main Defect": ("Text", "No", "Most frequent defect.", "Dropped stitch"),
    "Measurement Check": ("Text", "Yes", "Pass or Fail against the size spec.", "Pass"),
    "Spec Weight g": ("Integer", "No", "Target garment weight in grams.", "680"),
    "Avg Weight g": ("Integer", "No", "Measured average weight.", "684"),
    "Remarks": ("Text", "No", "Inspector comment.", ""),
}
FAC_COLS = {
    "Factory Code": ("Text", "Yes", "Short unique code used in every other file.", "GRL"),
    "Factory Name": ("Text", "Yes", "Legal name (fictional).", "Greyloom Knitwear Ltd."),
    "Area": ("Text", "Yes", "Industrial area.", "Konabari"),
    "District": ("Text", "Yes", "District.", "Gazipur"),
    "Address": ("Text", "No", "Postal address.", "Plot 14, Konabari, Gazipur"),
    "Supplier Since": ("Date", "No", "First Karbar order.", "01-Mar-2018"),
    "Gauges": ("Text", "Yes", "Gauges the factory can knit.", "5GG, 7GG, 12GG"),
    "Knitting M/C 3GG": ("Integer", "Yes", "Automatic flat knitting machines by gauge.", "0"),
    "Knitting M/C 5GG": ("Integer", "Yes", "", "60"),
    "Knitting M/C 7GG": ("Integer", "Yes", "", "180"),
    "Knitting M/C 12GG": ("Integer", "Yes", "", "90"),
    "Knitting M/C 14GG": ("Integer", "Yes", "", "0"),
    "Total Knitting M/C": ("Integer", "Yes", "Sum of the gauge columns.", "330"),
    "Linking M/C": ("Integer", "Yes", "Linking machines.", "160"),
    "Washing": ("Text", "Yes", "In-house or Subcontract.", "In-house"),
    "Knitting Shifts": ("Integer", "Yes", "Shifts per day in knitting.", "2"),
    "Knitting Hours/Day": ("Integer", "Yes", "Machine hours per day.", "20"),
    "Linking Hours/Day": ("Integer", "Yes", "Linking hours per day.", "10"),
    "Workforce": ("Integer", "No", "Total workers.", "2400"),
    "Certifications": ("Text", "No", "Social and environmental certificates.", "amfori BSCI, WRAP"),
    "amfori BSCI Rating": ("Text", "No", "Latest audit rating A-E.", "B"),
    "Last Social Audit": ("Date", "No", "Date of the latest social audit.", "14-May-2026"),
    "Factory Contact": ("Text", "No", "Main contact (fictional).", "Nazmul Haque"),
    "Contact Title": ("Text", "No", "", "GM - Merchandising"),
    "Karbar Merchandiser": ("Text", "Yes", "Karbar owner of the factory relationship.", "Farhana Rahman"),
    "Status": ("Text", "Yes", "Active or Inactive.", "Active"),
}
SHIP_COLS = {
    "Shipment ID": ("Text", "Yes", "Unique shipment reference.", "SH-2610-0412"),
    "PO No": ("Text", "Yes", "Links to the Order Book.", "71004521"),
    "Factory Code": ("Text", "Yes", "Shipping factory.", "GRL"),
    "Planned Ex-factory": ("Date", "Yes", "From the Order Book, repeated for audit.", "29-Oct-2026"),
    "Actual Ex-factory": ("Date", "Yes", "Date goods left the factory. On time if <= planned.", "29-Oct-2026"),
    "Shipped Qty": ("Integer", "Yes", "Pieces shipped.", "9600"),
    "Cartons": ("Integer", "Yes", "Carton count.", "534"),
    "Gross Weight kg": ("Decimal", "Yes", "Gross weight.", "7552.8"),
    "Ship Mode": ("Text", "Yes", "Sea, Air or Sea-Air.", "Sea"),
    "Port of Loading": ("Text", "Yes", "Chattogram or Dhaka (HSIA).", "Chattogram"),
    "Destination": ("Text", "Yes", "Karbar DC.", "Hamburg, Germany"),
    "ETD": ("Date", "Yes", "Estimated departure.", "03-Nov-2026"),
    "ETA": ("Date", "Yes", "Estimated arrival.", "07-Dec-2026"),
    "Vessel / Flight": ("Text", "No", "Vessel and voyage, or air waybill.", "MV Coral Horizon V.2611E"),
    "Invoice No": ("Text", "Yes", "Commercial invoice.", "KB-INV-2610-4412"),
    "Invoice Value USD": ("Decimal", "Yes", "Shipped Qty x FOB.", "158400.00"),
    "Air Freight Cost USD": ("Decimal", "No", "Only for Air shipments.", "45316.80"),
    "Freight Cost Borne By": ("Text", "Yes", "Karbar or Factory.", "Factory"),
    "Shipment Status": ("Text", "Yes", "At port, In transit or Delivered (as of the report date).", "In transit"),
    "Remarks": ("Text", "No", "", "Shipped 6 days late - air freight at factory cost"),
}


def write_book(path: Path, sheets, date_cols, about: str | None = None):
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for name, df, cols in sheets:
            df = df[[c for c in cols if c in df.columns]]
            df.to_excel(xw, sheet_name=name, index=False)
        dict_rows = []
        for name, df, cols in sheets:
            for c, (t, req, desc, ex) in cols.items():
                dict_rows.append({"Sheet": name, "Column": c, "Type": t, "Required": req, "Description": desc, "Example": ex})
        pd.DataFrame(dict_rows).to_excel(xw, sheet_name="Columns", index=False)
    style(path, date_cols)


HEADER_FILL = PatternFill("solid", fgColor="1F2A44")


def style(path: Path, date_cols):
    wb = load_workbook(path)
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        headers = [c.value for c in ws[1]]
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = HEADER_FILL
            c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[1].height = 30
        for i, h in enumerate(headers, 1):
            col = ws.cell(row=1, column=i).column_letter
            sample = [ws.cell(row=r, column=i).value for r in range(2, min(ws.max_row, 300) + 1)]
            width = max([len(str(h))] + [len(str(v)) for v in sample if v is not None]) + 2
            ws.column_dimensions[col].width = min(max(width, 9), 60 if ws.title == "Columns" else 42)
            fmt = None
            if h in date_cols or (isinstance(h, str) and ("Ex-factory" in h or h in ("ETD", "ETA", "Last Report"))):
                fmt = "dd-mmm-yyyy"
            elif isinstance(h, str) and ("USD" in h and "pc" not in h):
                fmt = "#,##0.00"
            elif isinstance(h, str) and (h.endswith(("Qty", "Cum", "Day")) or h in ("Order Qty", "Pcs")):
                fmt = "#,##0"
            elif isinstance(h, str) and ("Share" in h or "Probability" in h or "OTD" in h or "Pass 90d" in h):
                fmt = "0.0%"
            if fmt:
                for r in range(2, ws.max_row + 1):
                    ws.cell(row=r, column=i).number_format = fmt
        if ws.max_row > 1 and ws.title != "Columns":
            ws.auto_filter.ref = ws.dimensions
    wb.save(path)


def write_order_book(pos: list[PO]):
    head, lines = [], []
    for p in sorted(pos, key=lambda x: x.po_no):
        head.append({"PO No": p.po_no, "PO Date": p.po_date, "Season": p.season, "Department": p.dept,
                     "Style No": p.style_no, "Style Name": p.style_name, "Product Type": p.product, "Gauge": p.gauge,
                     "Yarn Composition": p.yarn, "Wash Required": p.wash, "Weight kg/pc": p.weight,
                     "Knitting Minutes/pc": p.knit_min, "Linking Std pcs/M/C/day": p.link_std,
                     "Factory Code": p.factory, "Factory Name": FAC[p.factory].name, "Order Qty": p.qty,
                     "FOB USD/pc": p.fob, "FOB Value USD": p.value, "Planned Ex-factory": p.exf,
                     "Planned Ship Mode": "Sea", "Port of Loading": "Chattogram", "Destination": p.destination,
                     "Delivery Terms": "FOB Chattogram", "Merchandiser": p.merch})
        sizes, ratios = SIZES[p.dept]
        r = random.Random(f"cc-{p.po_no}")
        for colour, cq in p.colours:
            row = {"PO No": p.po_no, "Colour": colour, "Colour Code": f"{colour[:3].upper()}-{r.randint(100, 999)}"}
            row.update({s: None for s in ALL_SIZES})
            for s, q in zip(sizes, split_qty(cq, ratios)):
                row[s] = q
            row["Line Qty"] = cq
            lines.append(row)
    write_book(OUT / "01_Order_Book.xlsx", [("PO Header", pd.DataFrame(head), PO_COLS), ("PO Lines", pd.DataFrame(lines), LINE_COLS)],
               ["PO Date", "Planned Ex-factory"])


def write_demo_uploads(demo_rows):
    for code in ("PBB", "HMR"):
        rows = [r for c, r in demo_rows if c == code]
        df = pd.DataFrame(rows)
        if code == "HMR" and len(df):
            df = df.copy()
            df.loc[df.index[0], "PO No"] = "71999999"                             # unknown PO
            if len(df) > 1:
                df.loc[df.index[1], "Linking Day"] = -40                          # negative quantity
            if len(df) > 2:
                df.loc[df.index[2], "Knitting Cum"] = int(df.iloc[2]["Knitting Cum"]) - 500   # cumulative went down
            name = f"DPR_{code}_{LAST_EXPECTED_REPORT:%Y-%m-%d}_with_errors.xlsx"
        else:
            name = f"DPR_{code}_{LAST_EXPECTED_REPORT:%Y-%m-%d}.xlsx"
        write_book(OUT / "demo_uploads" / name, [("Daily Production", df, DPR_COLS)], ["Report Date"])


def write_answer_key(report, preds, whatif):
    df = pd.DataFrame(preds).drop(columns=["_otd", "_pass90"]).sort_values(["Risk Score"], ascending=False)
    kpi = pd.DataFrame(report["kpis"], columns=["KPI", "Value"])
    params = pd.DataFrame([
        ("Today", f"{TODAY:%d %b %Y}"), ("Latest expected report date", f"{LAST_EXPECTED_REPORT:%d %b %Y}"),
        ("Weekend", "Friday"), ("Rate window", f"{RATE_WINDOW} working days, linear weights 1..n (latest heaviest)"),
        ("Supply-limited stage", f"work-in-front <= {FLOW_LIMIT_DAYS} days of its own rate, or not started: finishes at upstream finish + lag"),
        ("Stage lags (WD)", ", ".join(f"{k} {v}" for k, v in STAGE_LAG.items())),
        ("Yarn in-house to knitting start", f"{YARN_TO_KNIT_DAYS} days"),
        ("Air freight USD/kg", AIR_USD_PER_KG), ("Sea freight USD/kg", SEA_USD_PER_KG),
        ("Knitting machine minutes/day", KNIT_MC_MINUTES_PER_DAY),
        ("Score: schedule (max 50)", "slip>0: 35 + min(15, 1.5 x slip days); slack <=2 WD: 25; slack 3-5 WD: 12; else 0"),
        ("Score: T&A (max 20)", "yarn in-house late >=8 d: 12, 4-7 d: 8, 1-3 d: 4; PP sample approval late >=3 d: +8"),
        ("Score: factory OTD (max 15)", "min(15, (1 - OTD last 12 months) x 50)"),
        ("Score: quality (max 10)", "PO's latest inspection failed: 10; else factory AQL pass rate (90 d) < 85%: 5"),
        ("Score: freshness (max 5)", "in production only; working days without a report: 1 -> 3, 2+ -> 5"),
        ("Bands", "<25 On track, 25-49 Watch, 50-74 At risk, >=75 Critical; slip>0 -> at least At risk; slip>=7 -> Critical; "
                  "planned ex-factory < today and not shipped -> Late"),
        ("On-time probability", "share of the factory's last-12-month shipments whose slip <= predicted slack (calendar days), clamped 2%-98%"),
    ], columns=["Parameter", "Value"])
    sheets = [("KPIs", kpi, {"KPI": ("", "", "", ""), "Value": ("", "", "", "")}),
              ("Open PO Predictions", df, {c: ("", "", "", "") for c in df.columns}),
              ("Factory Summary", pd.DataFrame(report["factories"]), {c: ("", "", "", "") for c in report["factories"][0]}),
              ("Weekly Outlook", pd.DataFrame(report["weeks"]), {c: ("", "", "", "") for c in report["weeks"][0]}),
              ("Hero What-if", pd.DataFrame(whatif), {c: ("", "", "", "") for c in whatif[0]}),
              ("Parameters", params, {"Parameter": ("", "", "", ""), "Value": ("", "", "", "")})]
    path = OUT / "00_Answer_Key.xlsx"
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for name, d, _ in sheets:
            d.to_excel(xw, sheet_name=name, index=False)
    style(path, ["Planned Ex-factory", "Projected Ex-factory", "Last Report"])


def print_summary(report, whatif, n_dpr, n_ta, n_insp, n_ship, n_po):
    print(f"POs {n_po}, T&A rows {n_ta}, DPR rows {n_dpr}, inspections {n_insp}, shipments {n_ship}")
    for k, v in report["kpis"]:
        print(f"  {k}: {v}")
    df = report["df"]
    hero = df[df["Seed Scenario"] == "hero"].iloc[0]
    print("  HERO:", {k: hero[k] for k in ["PO No", "Bottleneck Stage", "Bottleneck Rate pcs/day", "Required Rate pcs/day",
                                          "Projected Ex-factory", "Slip Days", "Risk Score", "Band", "Drivers"]})
    for w in whatif:
        print("  WHAT-IF:", w)
    mism = df[((df["Seed Scenario"].isin(["ok", "watch", "donor"])) & df["Band"].isin(["At risk", "Critical", "Late"])) |
              ((df["Seed Scenario"].isin(["pred_late", "tight", "hero"])) & ~df["Band"].isin(["At risk", "Critical"]))]
    print(f"  Solver misses: {SOLVE_MISSES}")
    print(f"  Scenario/band mismatches: {len(mism)}")
    if len(mism):
        print(mism[["PO No", "Seed Scenario", "State", "Slip Days", "Slack WD", "Risk Score", "Band"]].to_string())


if __name__ == "__main__":
    main()
